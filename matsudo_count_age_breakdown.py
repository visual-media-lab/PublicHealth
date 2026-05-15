from openpyxl import load_workbook
import csv
from datetime import datetime

# Input and output files
INPUT_FILE = 'matsudo.xlsx'
OUTPUT_FILE = 'death_mat20_64_age_breakdown.csv'

# Five-year birth-year bands corresponding to ages 20-64.
# Keep this order for CSV and console output.
AGE_BANDS = [
    '2001～2005',
    '1996～2000',
    '1991～1995',
    '1986～1990',
    '1981～1985',
    '1976～1980',
    '1971～1975',
    '1966～1970',
    '1961～1965',
]
TARGET_AGE_BANDS = set(AGE_BANDS)

# Main aggregation categories
TARGETS = [
    {
        'dose': 4,
        'label': 'dose4_on_or_before_2022-09-19',
        'start_date': None,
        'end_date': datetime.strptime('2022-09-19', '%Y-%m-%d'),
    },
    {
        'dose': 4,
        'label': 'dose4_on_or_after_2022-09-20',
        'start_date': datetime.strptime('2022-09-20', '%Y-%m-%d'),
        'end_date': None,
    },
    {
        'dose': 5,
        'label': 'dose5_on_or_before_2023-09-19',
        'start_date': None,
        'end_date': datetime.strptime('2023-09-19', '%Y-%m-%d'),
    },
    {
        'dose': 5,
        'label': 'dose5_on_or_after_2023-09-20',
        'start_date': datetime.strptime('2023-09-20', '%Y-%m-%d'),
        'end_date': None,
    },
    {
        'dose': 6,
        'label': 'dose6_all_dates',
        'start_date': None,
        'end_date': None,
    },
    {
        'dose': 7,
        'label': 'dose7_all_dates',
        'start_date': None,
        'end_date': None,
    },
]

DATE_2022_09_19 = datetime.strptime('2022-09-19', '%Y-%m-%d')
DATE_2023_09_19 = datetime.strptime('2023-09-19', '%Y-%m-%d')

MAIN_METRICS = [
    'total_vaccinated',
    'death_within_180',
    'death_within_180_excluding_next_dose',
    'death_181_360',
    'death_181_360_excluding_next_dose',
]

EXTRA_METRICS = [
    'dose4_on_or_before_2022-09-19_without_dose5_within_360_days',
    'dose5_on_or_before_2023-09-19_without_dose6_within_360_days',
    'dose6_all_dates_without_dose7_within_360_days',
]


def is_yyyymmdd(value):
    if value is None:
        return False
    text = str(value).strip()
    return len(text) == 8 and text.isdigit()


def parse_yyyymmdd(value):
    return datetime.strptime(str(value).strip(), '%Y%m%d')


def get_dose_date(row, dose_number):
    """Return the vaccination date for the given dose number.

    The original worksheet layout stores dose N at column index:
    4 + 3 * (N - 1), using zero-based indexing.
    """
    index = 4 + 3 * (dose_number - 1)
    if index >= len(row):
        return None
    return parse_yyyymmdd(row[index]) if is_yyyymmdd(row[index]) else None


def get_death_date(row):
    if len(row) <= 3:
        return None
    return parse_yyyymmdd(row[3]) if is_yyyymmdd(row[3]) else None


def in_date_range(target_date, start_date=None, end_date=None):
    if target_date is None:
        return False
    if start_date is not None and target_date < start_date:
        return False
    if end_date is not None and target_date > end_date:
        return False
    return True


def no_next_dose_within_days(dose_date, next_dose_date, days=360):
    """Return True when the next dose was not given within the given number of days."""
    if dose_date is None:
        return False
    if next_dose_date is None:
        return True
    return (next_dose_date - dose_date).days > days


def new_metric_counts():
    return {metric: 0 for metric in MAIN_METRICS}


def init_result(label, dose):
    return {
        'label': label,
        'dose': dose,
        'all': new_metric_counts(),
        'by_age_band': {age_band: new_metric_counts() for age_band in AGE_BANDS},
    }


def increment_main_metric(result, age_band, metric):
    result['all'][metric] += 1
    result['by_age_band'][age_band][metric] += 1


def update_death_counts(result, age_band, death_date, dose_date, next_dose_date):
    if death_date is None or dose_date is None:
        return

    days_to_death = (death_date - dose_date).days
    has_no_next_dose = next_dose_date is None

    if 0 <= days_to_death <= 180:
        increment_main_metric(result, age_band, 'death_within_180')
        if has_no_next_dose:
            increment_main_metric(result, age_band, 'death_within_180_excluding_next_dose')
    elif 181 <= days_to_death <= 360:
        increment_main_metric(result, age_band, 'death_181_360')
        if has_no_next_dose:
            increment_main_metric(result, age_band, 'death_181_360_excluding_next_dose')


def init_extra_counts():
    return {
        metric: {
            'all': 0,
            'by_age_band': {age_band: 0 for age_band in AGE_BANDS},
        }
        for metric in EXTRA_METRICS
    }


def increment_extra_metric(extra_counts, metric, age_band):
    extra_counts[metric]['all'] += 1
    extra_counts[metric]['by_age_band'][age_band] += 1


def write_results(results, extra_counts):
    with open(OUTPUT_FILE, 'w', newline='', encoding='utf-8-sig') as csv_file:
        writer = csv.writer(csv_file)

        writer.writerow(['main_results'])
        writer.writerow([
            'category',
            'dose',
            'age_band',
            'total_vaccinated',
            'deaths_within_180_days',
            'deaths_within_180_days_excluding_next_dose_recipients',
            'deaths_181_to_360_days',
            'deaths_181_to_360_days_excluding_next_dose_recipients',
        ])

        for result in results:
            write_main_result_row(writer, result, 'all', result['all'])
            for age_band in AGE_BANDS:
                write_main_result_row(writer, result, age_band, result['by_age_band'][age_band])

        writer.writerow([])
        writer.writerow(['extra_metrics'])
        writer.writerow(['metric', 'age_band', 'count'])
        for metric in EXTRA_METRICS:
            writer.writerow([metric, 'all', extra_counts[metric]['all']])
            for age_band in AGE_BANDS:
                writer.writerow([metric, age_band, extra_counts[metric]['by_age_band'][age_band]])


def write_main_result_row(writer, result, age_band, counts):
    writer.writerow([
        result['label'],
        result['dose'],
        age_band,
        counts['total_vaccinated'],
        counts['death_within_180'],
        counts['death_within_180_excluding_next_dose'],
        counts['death_181_360'],
        counts['death_181_360_excluding_next_dose'],
    ])


def print_results(results, extra_counts):
    print('=== Main results ===')
    for result in results:
        print_main_result(result, 'all', result['all'])
        for age_band in AGE_BANDS:
            print_main_result(result, age_band, result['by_age_band'][age_band], indent=True)

    print('=== Extra metrics ===')
    for metric in EXTRA_METRICS:
        print(f'{metric}: all={extra_counts[metric]["all"]}')
        for age_band in AGE_BANDS:
            print(f'  {age_band}: {extra_counts[metric]["by_age_band"][age_band]}')


def print_main_result(result, age_band, counts, indent=False):
    prefix = '  ' if indent else ''
    print(
        f"{prefix}{result['label']} [{age_band}]: "
        f"total_vaccinated={counts['total_vaccinated']}, "
        f"deaths_within_180_days={counts['death_within_180']}, "
        f"deaths_within_180_days_excluding_next_dose_recipients="
        f"{counts['death_within_180_excluding_next_dose']}, "
        f"deaths_181_to_360_days={counts['death_181_360']}, "
        f"deaths_181_to_360_days_excluding_next_dose_recipients="
        f"{counts['death_181_360_excluding_next_dose']}"
    )


def main():
    workbook = load_workbook(INPUT_FILE, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))

    results = [init_result(target['label'], target['dose']) for target in TARGETS]
    extra_counts = init_extra_counts()

    for raw_row in rows:
        row = [str(cell).strip() if cell is not None else '' for cell in raw_row]
        if not row:
            continue

        age_band = row[0] if len(row) > 0 else ''
        if age_band not in TARGET_AGE_BANDS:
            continue

        death_date = get_death_date(row)
        dose_dates = {dose: get_dose_date(row, dose) for dose in range(4, 9)}

        for result, target in zip(results, TARGETS):
            dose_number = target['dose']
            dose_date = dose_dates.get(dose_number)
            next_dose_date = dose_dates.get(dose_number + 1)

            if not in_date_range(dose_date, target['start_date'], target['end_date']):
                continue

            increment_main_metric(result, age_band, 'total_vaccinated')
            update_death_counts(result, age_band, death_date, dose_date, next_dose_date)

        if (
            in_date_range(dose_dates[4], None, DATE_2022_09_19)
            and no_next_dose_within_days(dose_dates[4], dose_dates[5], 360)
        ):
            increment_extra_metric(
                extra_counts,
                'dose4_on_or_before_2022-09-19_without_dose5_within_360_days',
                age_band,
            )

        if (
            in_date_range(dose_dates[5], None, DATE_2023_09_19)
            and no_next_dose_within_days(dose_dates[5], dose_dates[6], 360)
        ):
            increment_extra_metric(
                extra_counts,
                'dose5_on_or_before_2023-09-19_without_dose6_within_360_days',
                age_band,
            )

        if dose_dates[6] is not None and no_next_dose_within_days(dose_dates[6], dose_dates[7], 360):
            increment_extra_metric(
                extra_counts,
                'dose6_all_dates_without_dose7_within_360_days',
                age_band,
            )

    write_results(results, extra_counts)
    print_results(results, extra_counts)
    print(f'CSV written to: {OUTPUT_FILE}')


if __name__ == '__main__':
    main()
