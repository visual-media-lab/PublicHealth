from openpyxl import load_workbook
import csv
from datetime import datetime

# Input and output files
INPUT_FILE = 'matsudo.xlsx'
OUTPUT_FILE = 'death_mat20_64_clean.csv'

# Five-year birth-year bands corresponding to ages 20-64
TARGET_AGE_BANDS = {
    '2001～2005', '1996～2000', '1991～1995', '1986～1990', '1981～1985', '1976～1980',
    '1971～1975', '1966～1970', '1961～1965'
}

# Main aggregation categories
TARGETS = [
    {
        'dose': 4,
        'label': 'dose4_on_or_before_2022-09-19',
        'start_date': None,
        'end_date': datetime.strptime('2022-09-19', '%Y-%m-%d'),
    },
    {
        'dose': 4,
        'label': 'dose4_on_or_after_2022-09-20',
        'start_date': datetime.strptime('2022-09-20', '%Y-%m-%d'),
        'end_date': None,
    },
    {
        'dose': 5,
        'label': 'dose5_on_or_before_2023-09-19',
        'start_date': None,
        'end_date': datetime.strptime('2023-09-19', '%Y-%m-%d'),
    },
    {
        'dose': 5,
        'label': 'dose5_on_or_after_2023-09-20',
        'start_date': datetime.strptime('2023-09-20', '%Y-%m-%d'),
        'end_date': None,
    },
    {
        'dose': 6,
        'label': 'dose6_all_dates',
        'start_date': None,
        'end_date': None,
    },
    {
        'dose': 7,
        'label': 'dose7_all_dates',
        'start_date': None,
        'end_date': None,
    },
]

DATE_2022_09_19 = datetime.strptime('2022-09-19', '%Y-%m-%d')
DATE_2023_09_19 = datetime.strptime('2023-09-19', '%Y-%m-%d')


def is_yyyymmdd(value):
    if value is None:
        return False
    text = str(value).strip()
    return len(text) == 8 and text.isdigit()


def parse_yyyymmdd(value):
    return datetime.strptime(str(value).strip(), '%Y%m%d')


def get_dose_date(row, dose_number):
    """Return the vaccination date for the given dose number.

    The original worksheet layout stores dose N at column index:
    4 + 3 * (N - 1), using zero-based indexing.
    """
    index = 4 + 3 * (dose_number - 1)
    if index >= len(row):
        return None
    return parse_yyyymmdd(row[index]) if is_yyyymmdd(row[index]) else None


def get_death_date(row):
    if len(row) <= 3:
        return None
    return parse_yyyymmdd(row[3]) if is_yyyymmdd(row[3]) else None


def in_date_range(target_date, start_date=None, end_date=None):
    if target_date is None:
        return False
    if start_date is not None and target_date < start_date:
        return False
    if end_date is not None and target_date > end_date:
        return False
    return True


def no_next_dose_within_days(dose_date, next_dose_date, days=360):
    """Return True when the next dose was not given within the given number of days."""
    if dose_date is None:
        return False
    if next_dose_date is None:
        return True
    return (next_dose_date - dose_date).days > days


def init_result(label, dose):
    return {
        'label': label,
        'dose': dose,
        'total_vaccinated': 0,
        'death_within_180': 0,
        'death_within_180_excluding_next_dose': 0,
        'death_181_360': 0,
        'death_181_360_excluding_next_dose': 0,
    }


def update_death_counts(result, death_date, dose_date, next_dose_date):
    if death_date is None or dose_date is None:
        return

    days_to_death = (death_date - dose_date).days
    has_no_next_dose = next_dose_date is None

    if 0 <= days_to_death <= 180:
        result['death_within_180'] += 1
        if has_no_next_dose:
            result['death_within_180_excluding_next_dose'] += 1
    elif 181 <= days_to_death <= 360:
        result['death_181_360'] += 1
        if has_no_next_dose:
            result['death_181_360_excluding_next_dose'] += 1


def write_results(results, extra_counts):
    with open(OUTPUT_FILE, 'w', newline='', encoding='utf-8-sig') as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow([
            'category',
            'dose',
            'total_vaccinated',
            'deaths_within_180_days',
            'deaths_within_180_days_excluding_next_dose_recipients',
            'deaths_181_to_360_days',
            'deaths_181_to_360_days_excluding_next_dose_recipients',
        ])

        for result in results:
            writer.writerow([
                result['label'],
                result['dose'],
                result['total_vaccinated'],
                result['death_within_180'],
                result['death_within_180_excluding_next_dose'],
                result['death_181_360'],
                result['death_181_360_excluding_next_dose'],
            ])

        writer.writerow([])
        writer.writerow(['extra_metric', 'count'])
        for label, count in extra_counts:
            writer.writerow([label, count])


def print_results(results, extra_counts):
    print('=== Main results ===')
    for result in results:
        print(
            f"{result['label']}: "
            f"total_vaccinated={result['total_vaccinated']}, "
            f"deaths_within_180_days={result['death_within_180']}, "
            f"deaths_within_180_days_excluding_next_dose_recipients="
            f"{result['death_within_180_excluding_next_dose']}, "
            f"deaths_181_to_360_days={result['death_181_360']}, "
            f"deaths_181_to_360_days_excluding_next_dose_recipients="
            f"{result['death_181_360_excluding_next_dose']}"
        )

    print('=== Extra metrics ===')
    for label, count in extra_counts:
        print(f'{label}: {count}')


def main():
    workbook = load_workbook(INPUT_FILE, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))

    results = [init_result(target['label'], target['dose']) for target in TARGETS]

    no_dose5_within_360_after_dose4_before_20220919 = 0
    no_dose6_within_360_after_dose5_before_20230919 = 0
    no_dose7_within_360_after_dose6_all = 0

    for raw_row in rows:
        row = [str(cell).strip() if cell is not None else '' for cell in raw_row]
        if not row:
            continue

        age_band = row[0] if len(row) > 0 else ''
        if age_band not in TARGET_AGE_BANDS:
            continue

        death_date = get_death_date(row)
        dose_dates = {dose: get_dose_date(row, dose) for dose in range(4, 9)}

        for result, target in zip(results, TARGETS):
            dose_number = target['dose']
            dose_date = dose_dates.get(dose_number)
            next_dose_date = dose_dates.get(dose_number + 1)

            if not in_date_range(dose_date, target['start_date'], target['end_date']):
                continue

            result['total_vaccinated'] += 1
            update_death_counts(result, death_date, dose_date, next_dose_date)

        if (
            in_date_range(dose_dates[4], None, DATE_2022_09_19)
            and no_next_dose_within_days(dose_dates[4], dose_dates[5], 360)
        ):
            no_dose5_within_360_after_dose4_before_20220919 += 1

        if (
            in_date_range(dose_dates[5], None, DATE_2023_09_19)
            and no_next_dose_within_days(dose_dates[5], dose_dates[6], 360)
        ):
            no_dose6_within_360_after_dose5_before_20230919 += 1

        if dose_dates[6] is not None and no_next_dose_within_days(dose_dates[6], dose_dates[7], 360):
            no_dose7_within_360_after_dose6_all += 1

    extra_counts = [
        (
            'dose4_on_or_before_2022-09-19_without_dose5_within_360_days',
            no_dose5_within_360_after_dose4_before_20220919,
        ),
        (
            'dose5_on_or_before_2023-09-19_without_dose6_within_360_days',
            no_dose6_within_360_after_dose5_before_20230919,
        ),
        (
            'dose6_all_dates_without_dose7_within_360_days',
            no_dose7_within_360_after_dose6_all,
        ),
    ]

    write_results(results, extra_counts)
    print_results(results, extra_counts)
    print(f'CSV written to: {OUTPUT_FILE}')


if __name__ == '__main__':
    main()
